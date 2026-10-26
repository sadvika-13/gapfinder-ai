
import json
import re
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path

from docx import Document
from fastapi import Depends, FastAPI, File, Form, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from pypdf import PdfReader
from sqlalchemy import (
    Column, DateTime, ForeignKey, Integer, JSON, String, Text,
    create_engine,
)
from sqlalchemy.orm import Session, declarative_base, sessionmaker


# --------------------------------------------------
# DATABASE — existing gapfinder.db is preserved
# --------------------------------------------------

DATABASE_URL = "sqlite:///./gapfinder.db"
MAX_FILE_SIZE = 10 * 1024 * 1024

engine = create_engine(
    DATABASE_URL,
    connect_args={"check_same_thread": False},
)
SessionLocal = sessionmaker(
    autocommit=False, autoflush=False, bind=engine
)
Base = declarative_base()


class SubjectAnalysis(Base):
    __tablename__ = "subject_analyses"

    id = Column(Integer, primary_key=True, index=True)
    student_name = Column(String, default="Student")
    subject = Column(String, default="General")
    filename = Column(String)
    file_type = Column(String)
    extracted_characters = Column(Integer, default=0)
    topics = Column(JSON, default=list)
    potential_gaps = Column(JSON, default=list)
    roadmap = Column(JSON, default=list)
    questions = Column(JSON, default=list)
    created_at = Column(
        DateTime, default=lambda: datetime.now(timezone.utc)
    )


class AnswerEvaluation(Base):
    __tablename__ = "answer_evaluations"

    id = Column(Integer, primary_key=True, index=True)
    analysis_id = Column(
        Integer, ForeignKey("subject_analyses.id"), index=True
    )
    topic = Column(String)
    question = Column(Text)
    answer = Column(Text)
    matched_concepts = Column(JSON, default=list)
    score = Column(Integer, default=0)
    feedback = Column(Text)
    is_gap = Column(Integer, default=1)
    created_at = Column(
        DateTime, default=lambda: datetime.now(timezone.utc)
    )


Base.metadata.create_all(bind=engine)

app = FastAPI(title="GapFinder AI - Subject File Analyzer")

app.add_middleware(
    CORSMiddleware,
    allow_origins=[
        "http://localhost:3000",
        "http://127.0.0.1:3000",
        "http://localhost:3001",
        "http://127.0.0.1:3001",
    ],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


def get_db():
    db = SessionLocal()
    try:
        yield db
    finally:
        db.close()


# --------------------------------------------------
# TOPIC KEYWORDS
# --------------------------------------------------

TOPIC_KEYWORDS = {
    "Machine Learning": [
        "machine learning", "supervised learning",
        "unsupervised learning", "regression", "classification",
        "clustering", "decision tree", "random forest",
        "overfitting", "underfitting", "model evaluation",
        "train test split", "neural network", "feature engineering",
    ],
    "Python": [
        "python", "variable", "data type", "list", "tuple",
        "dictionary", "function", "loop", "class",
        "exception handling", "module",
    ],
    "Data Structures": [
        "array", "linked list", "stack", "queue", "tree",
        "graph", "hash table", "sorting", "searching",
        "time complexity", "binary search", "recursion",
    ],
    "Database Management": [
        "database", "sql", "normalization", "primary key",
        "foreign key", "join", "transaction", "acid properties",
        "indexing", "entity relationship", "relational model",
    ],
    "Computer Networks": [
        "computer network", "tcp", "udp", "ip address",
        "dns", "http", "osi model", "network layer",
        "routing", "subnetting", "transmission control protocol",
    ],
    "Cybersecurity": [
        "cybersecurity", "encryption", "authentication",
        "authorization", "firewall", "malware", "phishing",
        "vulnerability", "siem", "access control",
        "network security",
    ],
    "Cloud Computing": [
        "cloud computing", "virtualization", "iaas", "paas",
        "saas", "public cloud", "private cloud", "hybrid cloud",
        "container", "scalability", "load balancing",
    ],
    "Natural Language Processing": [
        "natural language processing", "tokenization", "stemming",
        "lemmatization", "sentiment analysis",
        "named entity recognition", "text classification",
        "word embeddings", "language model",
    ],
    "Operating Systems": [
        "operating system", "process", "thread", "deadlock",
        "memory management", "paging", "scheduling",
        "semaphore", "virtual memory", "file system",
    ],
    "Software Engineering": [
        "software engineering", "software development life cycle",
        "sdlc", "agile", "scrum", "testing",
        "requirements analysis", "design pattern",
        "version control", "maintenance",
    ],
}


ETHICAL_HACKING_TOPICS = {
    "Ethical Hacking Fundamentals": [
        "ethical hacking", "ethical hacker", "penetration testing",
        "authorized testing", "permission", "scope of testing",
    ],
    "Vulnerability Assessment and Classification": [
        "vulnerability assessment", "vulnerability classification",
        "critical", "warning", "informational", "severity",
        "exposure", "risk assessment", "asset", "exploitation",
    ],
    "Penetration Testing Reports and Final Analysis": [
        "penetration testing report", "final analysis",
        "potential analysis", "final report", "findings",
        "executive summary", "final deliverable",
    ],
    "Report Structure and Presentation": [
        "overall structure", "report structure", "presentation",
        "business measurement", "process mapping", "limitations",
        "conclusion",
    ],
    "Remedial, Tactical and Strategic Remediation": [
        "remedial", "tactical", "strategic", "remediation",
        "short-term", "long-term",
    ],
    "Mitigation and Integration of Results": [
        "mitigation", "integration of results", "mitigation phases",
        "validate", "pilot", "implement", "defense planning",
    ],
    "Defense Planning and Architecture Review": [
        "defense planning", "architecture review",
        "security architecture", "centralization",
        "standardization", "technical review", "virtual architecture",
    ],
    "Security Awareness Training": [
        "security awareness", "awareness training",
        "security training", "role-based", "nist sp 800-50",
    ],
    "Incident Management and CERT": [
        "incident management", "incident response", "detect",
        "isolate", "eradicate", "recover",
        "computer emergency response team", "cert",
    ],
    "Security Policy and Access Control": [
        "security policy", "data classification",
        "organizational security", "access control",
        "separation of duties", "authentication", "authorization",
    ],
}


GENERIC_HEADINGS = {
    "final analysis", "potential analysis", "the document",
    "overall structure", "aligning findings", "business measurement",
    "presentation", "integration summary", "conclusion",
    "introduction", "references", "contents", "table of contents",
    "unit v", "unit 5", "unit-v",
}


# --------------------------------------------------
# MCQ QUESTION BANK
# Each question has 4 options, an answer and an explanation.
# --------------------------------------------------

MCQ_BANK = {
    "Ethical Hacking Fundamentals": [
        {
            "question": "What is the main purpose of ethical hacking?",
            "options": [
                "Steal confidential data",
                "Find security vulnerabilities with permission",
                "Damage computer systems",
                "Create malware",
            ],
            "correct_answer": "Find security vulnerabilities with permission",
            "explanation": "Ethical hackers test systems with authorization to identify security weaknesses.",
        },
        {
            "question": "What should an ethical hacker obtain before testing a system?",
            "options": [
                "Written authorization",
                "A personal social media account",
                "A public email address",
                "A software advertisement",
            ],
            "correct_answer": "Written authorization",
            "explanation": "Testing must be authorized and performed within the agreed scope.",
        },
    ],
    "Vulnerability Assessment and Classification": [
        {
            "question": "What is a vulnerability in cybersecurity?",
            "options": [
                "A system weakness",
                "A security update",
                "A backup file",
                "A strong password",
            ],
            "correct_answer": "A system weakness",
            "explanation": "A vulnerability is a weakness that may be exploited to affect a system.",
        },
        {
            "question": "Why are vulnerabilities classified by severity?",
            "options": [
                "To prioritize remediation",
                "To remove all backups",
                "To increase file size",
                "To disable all users",
            ],
            "correct_answer": "To prioritize remediation",
            "explanation": "Severity helps teams prioritize issues according to their potential impact and risk.",
        },
    ],
    "Penetration Testing Reports and Final Analysis": [
        {
            "question": "What is the purpose of a penetration testing report?",
            "options": [
                "Hide security issues",
                "Record findings and recommend fixes",
                "Delete system logs",
                "Create user accounts",
            ],
            "correct_answer": "Record findings and recommend fixes",
            "explanation": "The report documents vulnerabilities, their impact, evidence and remediation steps.",
        },
    ],
    "Report Structure and Presentation": [
        {
            "question": "Which item is important in a security assessment report?",
            "options": [
                "Findings and recommended actions",
                "Unrelated advertisements",
                "Personal entertainment preferences",
                "Unverified rumours",
            ],
            "correct_answer": "Findings and recommended actions",
            "explanation": "A useful report clearly communicates findings, evidence, impact and recommendations.",
        },
    ],
    "Remedial, Tactical and Strategic Remediation": [
        {
            "question": "What is the main goal of remediation?",
            "options": [
                "Fix identified security weaknesses",
                "Hide vulnerabilities from the team",
                "Remove security policies",
                "Ignore the assessment results",
            ],
            "correct_answer": "Fix identified security weaknesses",
            "explanation": "Remediation addresses identified weaknesses to reduce risk.",
        },
    ],
    "Mitigation and Integration of Results": [
        {
            "question": "Why should security fixes be validated after implementation?",
            "options": [
                "To confirm the issue has been addressed",
                "To create more vulnerabilities",
                "To avoid documenting changes",
                "To disable monitoring permanently",
            ],
            "correct_answer": "To confirm the issue has been addressed",
            "explanation": "Validation checks that the mitigation works and the weakness is no longer present.",
        },
    ],
    "Defense Planning and Architecture Review": [
        {
            "question": "What is a key goal of a security architecture review?",
            "options": [
                "Identify design weaknesses and improve protection",
                "Remove every access control",
                "Disable all security monitoring",
                "Avoid reviewing system design",
            ],
            "correct_answer": "Identify design weaknesses and improve protection",
            "explanation": "Architecture reviews assess security design and recommend improvements.",
        },
    ],
    "Security Awareness Training": [
        {
            "question": "What is a primary goal of security awareness training?",
            "options": [
                "Help users recognize and avoid security threats",
                "Teach users to share passwords",
                "Encourage opening unknown attachments",
                "Disable all security updates",
            ],
            "correct_answer": "Help users recognize and avoid security threats",
            "explanation": "Awareness training helps people recognize risks such as phishing and unsafe data handling.",
        },
    ],
    "Incident Management and CERT": [
        {
            "question": "What is an important early step in incident response?",
            "options": [
                "Identify and assess the incident",
                "Delete all evidence immediately",
                "Ignore security alerts",
                "Publish passwords publicly",
            ],
            "correct_answer": "Identify and assess the incident",
            "explanation": "Incident response starts with detection and analysis, followed by containment and recovery.",
        },
    ],
    "Security Policy and Access Control": [
        {
            "question": "What is the principle of least privilege?",
            "options": [
                "Give users only the access they need",
                "Give every user administrator access",
                "Share one password with everyone",
                "Disable authentication",
            ],
            "correct_answer": "Give users only the access they need",
            "explanation": "Least privilege limits access to the minimum required for a user's role.",
        },
    ],
    "Machine Learning": [
        {
            "question": "What is supervised learning?",
            "options": [
                "Learning from labelled examples",
                "Learning without any data",
                "Deleting the training dataset",
                "Only storing files",
            ],
            "correct_answer": "Learning from labelled examples",
            "explanation": "Supervised learning trains a model using examples with known target labels.",
        },
        {
            "question": "Why is a train-test split used?",
            "options": [
                "To evaluate performance on unseen data",
                "To duplicate every row",
                "To remove the target variable permanently",
                "To guarantee perfect accuracy",
            ],
            "correct_answer": "To evaluate performance on unseen data",
            "explanation": "A separate test set helps estimate how well a model generalizes to unseen examples.",
        },
    ],
    "Python": [
        {
            "question": "Which Python data type stores key-value pairs?",
            "options": ["List", "Dictionary", "Tuple", "String"],
            "correct_answer": "Dictionary",
            "explanation": "A Python dictionary stores values associated with keys.",
        },
        {
            "question": "Which keyword is used to define a function in Python?",
            "options": ["func", "define", "def", "function"],
            "correct_answer": "def",
            "explanation": "Python functions are declared using the def keyword.",
        },
    ],
    "Data Structures": [
        {
            "question": "Which data structure follows LIFO order?",
            "options": ["Queue", "Stack", "Graph", "Array"],
            "correct_answer": "Stack",
            "explanation": "A stack follows Last In, First Out order.",
        },
        {
            "question": "Which data structure normally follows FIFO order?",
            "options": ["Stack", "Queue", "Tree", "Heap"],
            "correct_answer": "Queue",
            "explanation": "A queue follows First In, First Out order.",
        },
    ],
    "Database Management": [
        {
            "question": "What is the main purpose of a primary key?",
            "options": [
                "Uniquely identify each row",
                "Store duplicate row identifiers",
                "Encrypt every database column",
                "Delete all records",
            ],
            "correct_answer": "Uniquely identify each row",
            "explanation": "A primary key uniquely identifies each record in a table.",
        },
        {
            "question": "What does SQL stand for?",
            "options": [
                "Structured Query Language",
                "Simple Question List",
                "System Quality Link",
                "Sequential Queue Logic",
            ],
            "correct_answer": "Structured Query Language",
            "explanation": "SQL is used to define, query and manipulate relational databases.",
        },
    ],
    "Computer Networks": [
        {
            "question": "Which protocol provides reliable, ordered data delivery?",
            "options": ["UDP", "TCP", "ARP only", "DHCP only"],
            "correct_answer": "TCP",
            "explanation": "TCP provides reliable, ordered delivery using acknowledgements and retransmission.",
        },
        {
            "question": "What is the primary purpose of DNS?",
            "options": [
                "Translate domain names to IP addresses",
                "Encrypt every local file",
                "Replace the operating system",
                "Create database tables",
            ],
            "correct_answer": "Translate domain names to IP addresses",
            "explanation": "DNS resolves domain names to corresponding network records, commonly IP addresses.",
        },
    ],
    "Cybersecurity": [
        {
            "question": "What is phishing?",
            "options": [
                "A deceptive attempt to steal information",
                "A routine data backup",
                "A method of cooling a computer",
                "A database indexing technique",
            ],
            "correct_answer": "A deceptive attempt to steal information",
            "explanation": "Phishing uses deceptive messages or sites to trick people into revealing information or taking unsafe actions.",
        },
        {
            "question": "What does a firewall primarily do?",
            "options": [
                "Filter network traffic according to rules",
                "Guarantee that no attack can occur",
                "Write application code automatically",
                "Replace all user authentication",
            ],
            "correct_answer": "Filter network traffic according to rules",
            "explanation": "Firewalls control permitted and blocked network traffic based on configured rules.",
        },
    ],
    "Cloud Computing": [
        {
            "question": "Which cloud service model provides virtual machines and storage?",
            "options": ["SaaS", "IaaS", "PaaS only", "DNS"],
            "correct_answer": "IaaS",
            "explanation": "Infrastructure as a Service provides computing infrastructure such as virtual machines and storage.",
        },
        {
            "question": "What is virtualization?",
            "options": [
                "Creating virtual versions of computing resources",
                "Deleting all operating systems",
                "Making every file public",
                "Turning off computer networks",
            ],
            "correct_answer": "Creating virtual versions of computing resources",
            "explanation": "Virtualization allows physical resources to host isolated virtual computing environments.",
        },
    ],
    "Natural Language Processing": [
        {
            "question": "What is tokenization in NLP?",
            "options": [
                "Splitting text into smaller units",
                "Encrypting a hard drive",
                "Sorting database tables",
                "Creating a network cable",
            ],
            "correct_answer": "Splitting text into smaller units",
            "explanation": "Tokenization divides text into units such as words, subwords or sentences.",
        },
        {
            "question": "What does sentiment analysis aim to identify?",
            "options": [
                "Opinion or emotional polarity in text",
                "The physical size of a computer",
                "The IP address of a website",
                "The number of database tables",
            ],
            "correct_answer": "Opinion or emotional polarity in text",
            "explanation": "Sentiment analysis estimates whether text expresses positive, negative or neutral sentiment.",
        },
    ],
    "Operating Systems": [
        {
            "question": "What is a process?",
            "options": [
                "A program in execution",
                "A type of network cable",
                "A database key",
                "A programming language comment",
            ],
            "correct_answer": "A program in execution",
            "explanation": "A process is an executing program together with its managed state and resources.",
        },
        {
            "question": "What is deadlock in an operating system?",
            "options": [
                "Processes wait indefinitely for resources held by one another",
                "A file is successfully saved",
                "A CPU finishes every task instantly",
                "A user changes a password",
            ],
            "correct_answer": "Processes wait indefinitely for resources held by one another",
            "explanation": "Deadlock occurs when processes are stuck waiting for resources in a circular dependency.",
        },
    ],
    "Software Engineering": [
        {
            "question": "What is the purpose of software testing?",
            "options": [
                "Find defects and check expected behaviour",
                "Guarantee software has no possible defects",
                "Remove all requirements",
                "Avoid checking the application",
            ],
            "correct_answer": "Find defects and check expected behaviour",
            "explanation": "Testing evaluates software behaviour and helps uncover defects and requirement mismatches.",
        },
        {
            "question": "What does SDLC stand for?",
            "options": [
                "Software Development Life Cycle",
                "System Data Link Control",
                "Simple Design Logic Code",
                "Software Download List Cache",
            ],
            "correct_answer": "Software Development Life Cycle",
            "explanation": "SDLC describes stages involved in planning, building, testing and maintaining software.",
        },
    ],
}


# --------------------------------------------------
# TEXT EXTRACTION AND TOPIC DETECTION
# --------------------------------------------------

def get_matches(text: str, keywords: list) -> list:
    lower_text = text.lower()
    matches = []

    for keyword in keywords:
        pattern = (
            r"(?<![a-z0-9])"
            + re.escape(keyword.lower())
            + r"(?![a-z0-9])"
        )
        if re.search(pattern, lower_text):
            matches.append(keyword)

    return matches


def extract_pdf_text(content: bytes) -> str:
    try:
        reader = PdfReader(BytesIO(content))

        if reader.is_encrypted:
            raise HTTPException(
                status_code=400,
                detail="This PDF is password-protected. Upload an unlocked PDF.",
            )

        return "\n".join(
            page.extract_text() or ""
            for page in reader.pages[:200]
        )

    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(
            status_code=400,
            detail="Could not read this PDF. Please upload a valid PDF file.",
        ) from exc


def extract_docx_text(content: bytes) -> str:
    try:
        document = Document(BytesIO(content))
        parts = [p.text for p in document.paragraphs]

        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))

        return "\n".join(parts)

    except Exception as exc:
        raise HTTPException(
            status_code=400,
            detail="Could not read this DOCX file. Please upload a valid Word document.",
        ) from exc


def clean_text(text: str) -> str:
    text = text.replace("\x00", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def detect_topics(text: str, subject: str = "") -> list:
    subject_lower = subject.lower()
    is_ethical_hacking = (
        "ethical hacking" in subject_lower
        or "penetration testing" in subject_lower
    )

    topic_dictionary = (
        ETHICAL_HACKING_TOPICS if is_ethical_hacking
        else TOPIC_KEYWORDS
    )

    topics = []

    for topic_name, keywords in topic_dictionary.items():
        matched = get_matches(text, keywords)

        if matched:
            topics.append({
                "name": topic_name,
                "matched_keywords": matched[:8],
                "keyword_count": len(matched),
            })

    if not is_ethical_hacking:
        headings = []

        for line in text.splitlines():
            line = line.strip()
            normalized = re.sub(r"^[\d.\-*\s]+", "", line).strip()
            key = normalized.lower().rstrip(":.")

            if (
                4 <= len(normalized) <= 80
                and len(normalized.split()) <= 10
                and key not in GENERIC_HEADINGS
                and (
                    re.match(r"^(unit|chapter|module|topic)\s+", key)
                    or re.match(r"^\d+[\.)]\s+", line)
                )
            ):
                headings.append(normalized.rstrip(":."))

        existing = {item["name"].lower() for item in topics}
        seen = set()

        for heading in headings:
            key = heading.lower()

            if key not in existing and key not in seen:
                seen.add(key)
                topics.append({
                    "name": heading,
                    "matched_keywords": [],
                    "keyword_count": 0,
                })

    return topics[:25]


# --------------------------------------------------
# MCQ GENERATION
# --------------------------------------------------

def make_mcq(topic_name: str, topic_keywords: list) -> dict:
    """Create a clearly labelled fallback question for an unknown topic."""
    main_term = topic_keywords[0] if topic_keywords else topic_name

    return {
        "topic": topic_name,
        "question": (
            f"Which approach is most appropriate when learning about "
            f"{main_term}?"
        ),
        "options": [
            f"Understand the concept of {main_term} and practise relevant examples",
            "Ignore the topic and skip all practice",
            "Memorize unrelated information without understanding",
            "Avoid reviewing the subject material",
        ],
        "correct_answer": (
            f"Understand the concept of {main_term} and practise relevant examples"
        ),
        "explanation": (
            f"Review the definition, purpose and examples of {main_term} "
            "in your study material. This is a general study question, "
            "not a test of detailed subject facts."
        ),
        "type": "MCQ",
    }


def build_report(text: str, subject: str) -> dict:
    topics = detect_topics(text, subject)

    subject_lower = subject.lower()
    is_ethical_hacking = (
        "ethical hacking" in subject_lower
        or "penetration testing" in subject_lower
    )

    topic_dictionary = (
        ETHICAL_HACKING_TOPICS if is_ethical_hacking
        else TOPIC_KEYWORDS
    )

    potential_gaps = []

    for topic_name, keywords in topic_dictionary.items():
        matched = get_matches(text, keywords)

        if not matched:
            potential_gaps.append({
                "topic": topic_name,
                "reason": (
                    "The expected keywords for this topic were not clearly "
                    "found in the uploaded document. Check your syllabus "
                    "to confirm whether this topic needs review."
                ),
            })

    roadmap = []

    for index, topic in enumerate(topics[:7], start=1):
        roadmap.append({
            "step": index,
            "topic": topic["name"],
            "action": (
                f"Review {topic['name']}, write a short explanation "
                "in your own words, and answer two practice questions."
            ),
        })

    questions = []

    for topic in topics[:8]:
        topic_name = topic["name"]
        bank = MCQ_BANK.get(topic_name)

        if bank:
            for item in bank[:2]:
                questions.append({
                    "topic": topic_name,
                    **item,
                    "type": "MCQ",
                })
        else:
            questions.append(
                make_mcq(
                    topic_name,
                    topic_dictionary.get(topic_name, []),
                )
            )

    if not questions:
        questions.append({
            "topic": "Uploaded material",
            "question": "What should you do first when studying your notes?",
            "options": [
                "Identify the main concepts",
                "Ignore the notes",
                "Skip the important sections",
                "Read unrelated material",
            ],
            "correct_answer": "Identify the main concepts",
            "explanation": (
                "Identifying the main concepts helps organize revision."
            ),
            "type": "MCQ",
        })

    return {
        "subject": subject.strip() or "General",
        "topics": topics,
        "potential_gaps": potential_gaps,
        "roadmap": roadmap,
        "questions": questions,
        "summary": (
            f"The document contains approximately {len(text.split())} words. "
            f"The system identified {len(topics)} possible topics. "
            f"It found {len(potential_gaps)} possible review areas. "
            "Verify these areas against your syllabus."
        ),
        "analysis_method": "Rule-based subject-specific keyword analysis",
        "limitations": (
            "This version uses predefined keywords and MCQ banks, not a "
            "generative AI model. A missing keyword does not prove a topic "
            "is missing. Some unrecognized topics receive general study "
            "questions instead of detailed subject questions."
        ),
    }


# --------------------------------------------------
# SHORT-ANSWER EVALUATION — retained for old reports
# --------------------------------------------------

def evaluate_answer(topic: str, answer: str) -> dict:
    answer_lower = answer.strip().lower()

    if len(answer_lower) < 10:
        return {
            "score": 0,
            "matched_concepts": [],
            "feedback": (
                "Your answer is too short to assess. Add a definition, "
                "key concepts and an example where appropriate."
            ),
            "is_gap": True,
        }

    all_topics = {
        **TOPIC_KEYWORDS,
        **ETHICAL_HACKING_TOPICS,
    }

    keywords = all_topics.get(topic, [])

    if not keywords:
        return {
            "score": 0,
            "matched_concepts": [],
            "feedback": (
                "A topic-specific answer key is not available. "
                "Please review your answer manually."
            ),
            "is_gap": True,
        }

    matched = get_matches(answer_lower, keywords)
    score = round((len(matched) / len(keywords)) * 100)

    if len(matched) >= 2:
        feedback = (
            "Some relevant keywords appear in your answer. "
            "Check that your explanation is accurate and complete."
        )
        is_gap = False
    else:
        feedback = (
            "Few recognized keywords appear in your answer. Review the "
            "topic and include relevant concepts, definitions and examples."
        )
        is_gap = True

    return {
        "score": score,
        "matched_concepts": matched,
        "feedback": feedback,
        "is_gap": is_gap,
    }


# --------------------------------------------------
# HOME
# --------------------------------------------------

@app.get("/")
def home():
    return {
        "message": "GapFinder AI file analysis backend is running!"
    }


# --------------------------------------------------
# UPLOAD AND ANALYZE FILE
# --------------------------------------------------

@app.post("/analyze-file")
async def analyze_file(
    student_name: str = Form("Student"),
    subject: str = Form(""),
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
):
    filename = Path(file.filename or "").name
    extension = Path(filename).suffix.lower()

    if extension not in {".pdf", ".docx"}:
        raise HTTPException(
            status_code=400,
            detail="Only PDF and DOCX files are supported.",
        )

    content = await file.read(MAX_FILE_SIZE + 1)

    if len(content) > MAX_FILE_SIZE:
        raise HTTPException(
            status_code=413,
            detail="File is too large. Maximum allowed size is 10 MB.",
        )

    if not content:
        raise HTTPException(
            status_code=400,
            detail="The uploaded file is empty.",
        )

    if extension == ".pdf":
        extracted_text = extract_pdf_text(content)
        file_type = "PDF"
    else:
        extracted_text = extract_docx_text(content)
        file_type = "DOCX"

    extracted_text = clean_text(extracted_text)

    if len(extracted_text) < 30:
        raise HTTPException(
            status_code=400,
            detail=(
                "Very little readable text was found. "
                "For a scanned PDF, OCR support may be required."
            ),
        )

    report = build_report(extracted_text, subject)

    record = SubjectAnalysis(
        student_name=student_name.strip() or "Student",
        subject=subject.strip() or "General",
        filename=filename,
        file_type=file_type,
        extracted_characters=len(extracted_text),
        topics=report["topics"],
        potential_gaps=report["potential_gaps"],
        roadmap=report["roadmap"],
        questions=report["questions"],
    )

    db.add(record)
    db.commit()
    db.refresh(record)

    return {
        "id": record.id,
        "student_name": record.student_name,
        "filename": filename,
        "file_type": file_type,
        "extracted_characters": len(extracted_text),
        **report,
    }


# --------------------------------------------------
# EVALUATE MCQ AND SHORT ANSWERS
# --------------------------------------------------

@app.post("/evaluate-answers")
def evaluate_answers(
    analysis_id: int = Form(...),
    answers_json: str = Form(...),
    db: Session = Depends(get_db),
):
    try:
        answers = json.loads(answers_json)
    except json.JSONDecodeError as exc:
        raise HTTPException(
            status_code=400,
            detail="answers_json must contain valid JSON.",
        ) from exc

    if not isinstance(answers, list) or not answers:
        raise HTTPException(
            status_code=400,
            detail="Please provide a non-empty list of answers.",
        )

    record = db.query(SubjectAnalysis).filter(
        SubjectAnalysis.id == analysis_id
    ).first()

    if not record:
        raise HTTPException(
            status_code=404,
            detail="Analysis report not found.",
        )

    saved_questions = record.questions or []
    results = []
    gaps_by_topic = {}

    for item in answers:
        if not isinstance(item, dict):
            continue

        topic = str(item.get("topic", "")).strip()
        question_text = str(item.get("question", "")).strip()
        answer = str(item.get("answer", "")).strip()

        if not topic or not question_text:
            continue

        original = next(
            (
                q for q in saved_questions
                if q.get("topic") == topic
                and q.get("question") == question_text
            ),
            None,
        )

        if original and original.get("options"):
            correct_answer = str(
                original.get("correct_answer", "")
            ).strip()
            explanation = str(
                original.get("explanation", "")
            ).strip()

            is_correct = (
                bool(answer)
                and answer.casefold() == correct_answer.casefold()
            )

            score = 100 if is_correct else 0
            is_gap = not is_correct
            feedback = (
                "Correct! Well done."
                if is_correct
                else "Incorrect. Review the correct answer and explanation."
            )
            matched_concepts = [correct_answer] if is_correct else []

        else:
            evaluation = evaluate_answer(topic, answer)
            score = evaluation["score"]
            is_gap = evaluation["is_gap"]
            feedback = evaluation["feedback"]
            matched_concepts = evaluation["matched_concepts"]
            is_correct = None
            correct_answer = ""
            explanation = ""

        saved = AnswerEvaluation(
            analysis_id=analysis_id,
            topic=topic,
            question=question_text,
            answer=answer,
            matched_concepts=matched_concepts,
            score=score,
            feedback=feedback,
            is_gap=int(is_gap),
        )
        db.add(saved)

        result_item = {
            "topic": topic,
            "question": question_text,
            "score": score,
            "matched_concepts": matched_concepts,
            "feedback": feedback,
            "needs_review": is_gap,
            "is_gap": is_gap,
        }

        if is_correct is not None:
            result_item["is_correct"] = is_correct
            result_item["correct_answer"] = correct_answer
            result_item["explanation"] = explanation

        results.append(result_item)

        if is_gap:
            gaps_by_topic[topic] = {
                "topic": topic,
                "reason": feedback,
            }

    if not results:
        raise HTTPException(
            status_code=400,
            detail="No valid answers found.",
        )

    db.commit()

    mcq_results = [
        result for result in results if "is_correct" in result
    ]
    correct_count = sum(
        1 for result in mcq_results
        if result["is_correct"] is True
    )
    mcq_count = len(mcq_results)

    return {
        "analysis_id": analysis_id,
        "evaluated_count": len(results),
        "results": results,
        "correct_count": correct_count,
        "mcq_count": mcq_count,
        "percentage": (
            round(correct_count / mcq_count * 100)
            if mcq_count else 0
        ),
        "potential_gaps": list(gaps_by_topic.values()),
        "notice": (
            "MCQs are graded against the saved answer key. "
            "Older short-answer questions receive keyword-based feedback."
        ),
    }


# --------------------------------------------------
# ANALYSIS HISTORY
# --------------------------------------------------

@app.get("/analyses")
def get_analyses(db: Session = Depends(get_db)):
    records = (
        db.query(SubjectAnalysis)
        .order_by(SubjectAnalysis.id.desc())
        .limit(50)
        .all()
    )

    return [
        {
            "id": record.id,
            "student_name": record.student_name,
            "subject": record.subject,
            "filename": record.filename,
            "file_type": record.file_type,
            "created_at": (
                record.created_at.isoformat()
                if record.created_at else None
            ),
            "topics": record.topics or [],
            "potential_gaps": record.potential_gaps or [],
        }
        for record in records
    ]


@app.get("/analyses/{analysis_id}/evaluations")
def get_evaluations(
    analysis_id: int,
    db: Session = Depends(get_db),
):
    record = db.query(SubjectAnalysis).filter(
        SubjectAnalysis.id == analysis_id
    ).first()

    if not record:
        raise HTTPException(
            status_code=404,
            detail="Analysis report not found.",
        )

    evaluations = (
        db.query(AnswerEvaluation)
        .filter(AnswerEvaluation.analysis_id == analysis_id)
        .order_by(AnswerEvaluation.id.asc())
        .all()
    )

    return [
        {
            "id": item.id,
            "topic": item.topic,
            "question": item.question,
            "answer": item.answer,
            "matched_concepts": item.matched_concepts or [],
            "score": item.score,
            "feedback": item.feedback,
            "needs_review": bool(item.is_gap),
            "created_at": (
                item.created_at.isoformat()
                if item.created_at else None
            ),
        }
        for item in evaluations
    ]
